from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/ComplianceCopilot-Working-and-Button-Guide.md'
OUT = ROOT / 'output/user-guide/ComplianceCopilot-User-Guide.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.75), Inches(.7)
section.left_margin, section.right_margin = Inches(.7), Inches(.7)
section.header_distance, section.footer_distance = Inches(.3), Inches(.3)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.widow_control = True
    for border in list(st.element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(30)
doc.styles['Title'].paragraph_format.space_after = Pt(16)
doc.styles['Subtitle'].font.size = Pt(15)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(18)
for name, size in [('Heading 1', 17), ('Heading 2', 12.5), ('Heading 3', 11)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(17 if name == 'Heading 1' else 10)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = 'ComplianceCopilot   |   Workflow and button guide'
header.runs[0].font.name = 'Arial'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
footer = section.footer.paragraphs[0]
footer.text = '25 October 2026   •   App user guide'
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.35))
footer.add_run('\tPage ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for r in footer.runs:
    r.font.name = 'Arial'
    r.font.size = Pt(8)

def inline(p, text):
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        code = part.startswith('`') and part.endswith('`')
        r = p.add_run(part[2:-2] if bold else part[1:-1] if code else part)
        r.bold = bold

def make_table(lines):
    rows = [[v.strip() for v in l.strip().strip('|').split('|')] for l in lines]
    rows = [r for r in rows if not all(re.fullmatch(r'[:\- ]+', v) for v in r)]
    count = len(rows[0])
    widths = [1.6, 2.85, 2.65] if count == 3 else [1.4, 1.35, 2.15, 2.2] if count == 4 else [7.1 / count] * count
    table = doc.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    table._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = table.add_row()
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        pr.append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            repeat.set(qn('w:val'), 'true')
            pr.append(repeat)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '193553' if i == 0 else 'F3F6F9' if i % 2 else 'FFFFFF')
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, val in [('top', '90'), ('bottom', '90'), ('left', '105'), ('right', '105')]:
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), val)
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = i == 0
            inline(p, value)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else RGBColor(0, 0, 0)
                if i == 0:
                    run.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run().font.size = Pt(3)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
first_section = next(i for i, l in enumerate(lines) if l.startswith('## '))
doc.add_paragraph('ComplianceCopilot user guide', 'Title')
doc.add_paragraph('Detailed workflow and button reference', 'Subtitle')
doc.add_paragraph('For SME owners and product demonstrations\n25 October 2026')
for text in '\n'.join(lines[6:first_section]).strip().split('\n\n'):
    if text.strip():
        p = doc.add_paragraph()
        inline(p, text.replace('\n', ' '))
doc.add_paragraph('How to use this guide', 'Heading 2')
doc.add_paragraph('Read Sections 1 to 3 for the workflow. Use the contents to find a screen, then look for its button tables. Sections 14 to 18 explain growth, financial planning, scheduling, AI and calculators. Section 21 gives complete worked journeys.')
doc.add_page_break()
contents = doc.add_paragraph('Contents', 'TOC Heading')
contents.style.font.name = 'Arial'
contents.style.font.size = Pt(17)
contents.style.font.color.rgb = RGBColor(0, 0, 0)
contents.style.font.bold = True
outline = OxmlElement('w:outlineLvl')
outline.set(qn('w:val'), '9')
contents._p.get_or_add_pPr().append(outline)
toc = doc.add_paragraph()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'TOC \\o "1-1" \\h \\z \\u')
toc._p.append(field)
doc.add_page_break()

i = first_section
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        group = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            group.append(lines[i])
            i += 1
        make_table(group)
        continue
    if line.startswith('### '):
        doc.add_paragraph(line[4:], 'Heading 2')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    else:
        p = doc.add_paragraph()
        inline(p, line)
        if re.match(r'^\d+\. ', line):
            p.paragraph_format.left_indent = Inches(.16)
            p.paragraph_format.first_line_indent = Inches(-.16)
    i += 1

doc.core_properties.title = 'ComplianceCopilot user guide'
doc.core_properties.subject = 'Detailed workflow button reference and practical user examples'
doc.core_properties.author = 'ComplianceCopilot'
doc.core_properties.keywords = 'SME compliance user guide workflow buttons'
settings = doc.settings.element
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
settings.append(update)
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
print('Tables:', len(doc.tables))
